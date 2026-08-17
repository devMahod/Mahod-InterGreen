# -*- coding: utf-8 -*-
"""
Golden-data extraction + independent verification oracle.

Reads both completed IG_matrix workbooks (source evidence, read-only),
extracts every conflict row into JSON fixtures for the C# tests, and
INDEPENDENTLY recomputes the Legacy algorithm in Python, comparing
against the workbook's cached values.

This script is a development oracle only (execution prompt §3):
- not a production dependency
- output saved under tests/fixtures/ and test-results/
"""
import json, math, os, sys, warnings
import openpyxl

warnings.filterwarnings("ignore")

BASE = r"C:\Users\arthurf\Downloads\InterGreens\materials\Inter-green Automation"
EX1 = os.path.join(BASE, "03 Example 1", "05_PINES-HABAL- SHEM-TOM_IG_matrix_2026-07-28-maya.xlsx")
EX2 = os.path.join(BASE, "04 Example 2", "Abarbanel-HaMaccabim_IG_matrix_2026-07-05.xlsx")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIXTURES = os.path.join(ROOT, "tests", "fixtures")
RESULTS = os.path.join(ROOT, "test-results")
os.makedirs(FIXTURES, exist_ok=True)
os.makedirs(RESULTS, exist_ok=True)


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def load(path):
    wbF = openpyxl.load_workbook(path, data_only=False)
    wbV = openpyxl.load_workbook(path, data_only=True)
    return wbF, wbV


def detect_variant(wbF, wbV):
    """Detect template variant from the constants-block labels (col H) and row-1 indices."""
    ws = wbV["Parameters"]
    labels = {}
    for row in ws.iter_rows(min_row=2, max_row=16, min_col=8, max_col=9):
        h, i = row
        if h.value:
            labels[h.row] = str(h.value)
    # find which row holds vehicle length vs ped speed
    veh_row = next((r for r, t in labels.items() if "אורך רכב" in t), None)
    ped_row = next((r for r, t in labels.items() if 'מהירות ה"ר' in t or "מהירות ה" in t), None)
    rea_row = next((r for r, t in labels.items() if "זמן תגובה" in t), None)
    dec_row = next((r for r, t in labels.items() if "תאוט" in t), None)
    const = lambda r: ws.cell(row=r, column=9).value
    row1 = {c: ws.cell(row=1, column=c).value for c in range(1, 7)}
    e_header = ws.cell(row=2, column=5).value
    if veh_row is None:
        variant = "V1"  # per-movement vehicle length via VLOOKUP col E
        consts = dict(pedSpeed=const(ped_row), reaction=const(rea_row), decel=const(dec_row),
                      vehLenGlobal=None)
        inbar_flag = ws.cell(row=7, column=9).value
        inbar_len = ws.cell(row=8, column=9).value
        consts["inbarFlag"] = inbar_flag
        consts["inbarDefaultVehicleLength"] = inbar_len
    else:
        variant = "V2"  # global vehicle length in constants block
        consts = dict(pedSpeed=const(ped_row), reaction=const(rea_row), decel=const(dec_row),
                      vehLenGlobal=const(veh_row), inbarFlag=None, inbarDefaultVehicleLength=None)
    params = {}
    for row in ws.iter_rows(min_row=3, max_row=14, min_col=1, max_col=6):
        name = row[0].value
        if not name:
            continue
        params[str(name)] = {
            "interurban": row[1].value,
            "fastKph": row[2].value,
            "slowKph": row[3].value,
            "colE": row[4].value,   # V1: vehicle length; V2: additional length (0)
            "colF": row[5].value,   # V1: addition to Inbar distances
        }
    return variant, consts, params, {"row1": row1, "colEHeader": e_header}


def extract_rows(wbF, wbV):
    wsF, wsV = wbF["Input Distances"], wbV["Input Distances"]
    rows = []
    for r in range(3, wsF.max_row + 1):
        a = wsV.cell(row=r, column=1).value
        b = wsV.cell(row=r, column=2).value
        if a is None and b is None:
            continue
        if not is_num(a):
            continue
        raw = lambda col: wsF.cell(row=r, column=col).value
        cach = lambda col: wsV.cell(row=r, column=col).value
        # raw D..K: preserve blank-vs-zero. A formula cell counts as populated.
        def rawpt(col):
            v = raw(col)
            if v is None:
                return None
            if isinstance(v, str) and v.startswith("="):
                cv = cach(col)
                return cv if cv is not None else None
            return v
        row = {
            "conflictNo": a,
            "clearing": cach(2),
            "entering": cach(3),
            "points": [
                {"cd": rawpt(4), "ed": rawpt(5)},
                {"cd": rawpt(6), "ed": rawpt(7)},
                {"cd": rawpt(8), "ed": rawpt(9)},
                {"cd": rawpt(10), "ed": rawpt(11)},
            ],
            "cached": {
                "vFast": cach(12), "vSlow": cach(13), "vEnter": cach(14),
                "pointTimes": [
                    {"tFast": cach(16), "tSlow": cach(17), "tEnter": cach(18), "ig": cach(19)},
                    {"tFast": cach(20), "tSlow": cach(21), "tEnter": cach(22), "ig": cach(23)},
                    {"tFast": cach(24), "tSlow": cach(25), "tEnter": cach(26), "ig": cach(27)},
                    {"tFast": cach(28), "tSlow": cach(29), "tEnter": cach(30), "ig": cach(31)},
                ],
                "definingPoint": cach(32),
                "clearingSg": cach(33), "enteringSg": cach(34),
                "finalIg": cach(37),
                "manualRoundingFlag": cach(38),
            },
        }
        rows.append(row)
    return rows


# ---------------- independent Legacy reimplementation ----------------

def clearing_time(dist, v, veh_len, c):
    """Workbook column P/Q semantics. dist: blank==0. Returns -1 on math error (IFERROR)."""
    ped, reaction, decel = c["pedSpeed"], c["reaction"], c["decel"]
    d = 0.0 if dist is None else float(dist)
    try:
        if v == ped:
            return d / ped
        kph = v * 3.6
        coeff = -1.0 if (kph > 80 or v == ped or v == 0) else 1.5 - 1.5 * kph / 80.0
        if coeff == -1.0:
            return reaction + v / (2 * decel) + (d + veh_len) / v
        brake = v * v / (2 * decel)
        return (-v + math.sqrt(v * v + 2 * (brake + d + veh_len) * coeff)) / coeff + reaction
    except Exception:
        return -1.0


def resolve_speeds(clearing, entering, params, c):
    def lk(name, key):
        p = params.get(str(name))
        if p is None:
            return None
        v = p[key]
        return v if is_num(v) else None
    fast = lk(clearing, "fastKph")
    slow = lk(clearing, "slowKph")
    ent = lk(entering, "fastKph")
    vF = fast / 3.6 if fast is not None else c["pedSpeed"]
    vS = slow / 3.6 if slow is not None else c["pedSpeed"]
    vE = ent / 3.6 if ent is not None else 0.0
    return vF, vS, vE


def resolve_veh_len(variant, clearing, params, c):
    if variant == "V2":
        return c["vehLenGlobal"]
    p = params.get(str(clearing))
    if p is not None and is_num(p["colE"]):
        return p["colE"]
    return None  # only reached in ped branch, where it is unused


def compute_row(row, variant, params, c):
    vF, vS, vE = resolve_speeds(row["clearing"], row["entering"], params, c)
    veh_len = resolve_veh_len(variant, row["clearing"], params, c)
    igs = []      # S, W, AA, AE per point (None == workbook "")
    times = []
    for i, pt in enumerate(row["points"]):
        cd, ed = pt["cd"], pt["ed"]
        if i > 0 and cd is None:
            igs.append(None); times.append(None)
            continue
        tF = clearing_time(cd, vF, veh_len if veh_len is not None else 0.0, c)
        tS = clearing_time(cd, vS, veh_len if veh_len is not None else 0.0, c)
        tE = 0.0 if vE == 0 else (0.0 if ed is None else float(ed)) / vE
        ig = None if (vF == 0 and vS == 0) else max(tF, tS) - tE
        igs.append(ig)
        times.append({"tFast": tF, "tSlow": tS, "tEnter": tE})
    numeric = [x for x in igs if x is not None]
    raw_max = max(numeric) if numeric else None
    # defining point per AF
    defining = None
    if any(pt["cd"] is not None or pt["ed"] is not None for pt in row["points"]):
        vals = [x if x is not None else -math.inf for x in igs]
        m = max(vals)
        defining = "Point " + str(vals.index(m) + 1)
    # FINAL IG per AK
    final = None
    if raw_max is not None:
        i = int(math.floor(raw_max))
        if i != 0:
            frac = math.fmod(raw_max, i)
            final = i if frac < 0.1 else i + 1
    return {"vFast": vF, "vSlow": vS, "vEnter": vE, "vehLen": veh_len,
            "igs": igs, "times": times, "rawMax": raw_max,
            "defining": defining, "finalIg": final}


def close(a, b, tol=1e-9):
    if a is None and b is None:
        return True, 0.0
    if a is None or b is None:
        return False, None
    if not (is_num(a) and is_num(b)):
        return a == b, 0.0
    d = abs(a - b)
    return d <= tol, d


def verify(name, path):
    wbF, wbV = load(path)
    variant, consts, params, meta = detect_variant(wbF, wbV)
    rows = extract_rows(wbF, wbV)
    report = {"example": name, "file": os.path.basename(path), "variant": variant,
              "consts": consts, "rows": len(rows)}
    final_ok = final_bad = 0
    inter_ok = inter_bad = 0
    max_delta = 0.0
    mismatches = []
    ceil_delta_rows = []
    below3 = []
    multipoint = []
    missing_measurement = []
    for row in rows:
        comp = compute_row(row, variant, params, consts)
        exp = row["cached"]
        okF, _ = close(comp["finalIg"], exp["finalIg"] if is_num(exp["finalIg"]) else None)
        if okF:
            final_ok += 1
        else:
            final_bad += 1
            mismatches.append((row["conflictNo"], "FINAL", exp["finalIg"], comp["finalIg"]))
        # intermediates: per-point ig columns S/W/AA/AE
        row_inter_ok = True
        for i in range(4):
            expected_ig = exp["pointTimes"][i]["ig"]
            expected_ig = expected_ig if is_num(expected_ig) else None
            got = comp["igs"][i]
            ok, d = close(got, expected_ig, tol=1e-8)
            if not ok:
                row_inter_ok = False
                mismatches.append((row["conflictNo"], f"IG_P{i+1}", expected_ig, got))
            elif d is not None:
                max_delta = max(max_delta, d)
        if row_inter_ok:
            inter_ok += 1
        else:
            inter_bad += 1
        # rounding strategy delta: workbook vs plain ceiling
        if comp["rawMax"] is not None and is_num(exp["finalIg"]):
            ceil_v = math.ceil(comp["rawMax"] - 1e-12)
            if ceil_v != exp["finalIg"]:
                ceil_delta_rows.append((row["conflictNo"], f'{row["clearing"]}→{row["entering"]}',
                                       round(comp["rawMax"], 3), exp["finalIg"], ceil_v))
        if is_num(exp["finalIg"]) and exp["finalIg"] < 3:
            below3.append(row["conflictNo"])
        pts = sum(1 for pt in row["points"] if pt["cd"] is not None or pt["ed"] is not None)
        if pts > 1:
            multipoint.append((row["conflictNo"], f'{row["clearing"]}→{row["entering"]}', pts,
                               row["points"]))
        cd1, ed1 = row["points"][0]["cd"], row["points"][0]["ed"]
        if cd1 is None:
            kind = "NO_MEASUREMENT_AT_ALL" if ed1 is None else "MISSING_CLEARING_MEASUREMENT"
            missing_measurement.append((row["conflictNo"], f'{row["clearing"]}→{row["entering"]}',
                                        kind, exp["finalIg"]))
    report.update({
        "finalOk": final_ok, "finalBad": final_bad,
        "interOk": inter_ok, "interBad": inter_bad,
        "maxDelta": max_delta, "mismatches": mismatches[:20],
        "ceilDeltaRows": ceil_delta_rows, "below3": below3,
        "multipoint": [(m[0], m[1], m[2]) for m in multipoint],
        "multipointDetail": multipoint,
        "missingMeasurement": missing_measurement,
    })
    # fixture for C# tests
    fixture = {
        "schemaVersion": "1.0",
        "example": name,
        "sourceFile": os.path.basename(path),
        "variant": variant,
        "constants": consts,
        "parameters": params,
        "rows": rows,
    }
    fx_path = os.path.join(FIXTURES, f"golden-{name}.json")
    with open(fx_path, "w", encoding="utf-8") as f:
        json.dump(fixture, f, ensure_ascii=False, indent=1)
    return report


def main():
    r1 = verify("example1", EX1)
    r2 = verify("example2", EX2)
    lines = ["# GOLDEN_EXTRACTION_REPORT", "",
             "Independent Python oracle vs workbook cached values.", ""]
    for r in (r1, r2):
        lines += [f"## {r['example']} — {r['file']}", "",
                  f"- variant: **{r['variant']}**, constants: `{r['consts']}`",
                  f"- rows: **{r['rows']}**",
                  f"- FINAL IG match: **{r['finalOk']} / {r['rows']}**",
                  f"- intermediate (S/W/AA/AE) rows fully matching: **{r['interOk']} / {r['rows']}**",
                  f"- max abs delta on matched intermediates: **{r['maxDelta']:.3e}**",
                  f"- rows where FINAL IG < 3: **{r['below3'] or 'NONE'}**",
                  f"- multi-point rows: **{r['multipoint'] or 'NONE'}**",
                  f"- missing-measurement rows: **{r['missingMeasurement'] or 'NONE'}**",
                  f"- workbook-vs-ceiling delta rows ({len(r['ceilDeltaRows'])}):", ""]
        for c in r["ceilDeltaRows"]:
            lines.append(f"  - conflict {c[0]} {c[1]}: raw={c[2]} workbook={c[3]} ceil={c[4]}")
        if r["mismatches"]:
            lines.append(f"- MISMATCHES (first 20): {r['mismatches']}")
        lines.append("")
    total_ok = r1["finalOk"] + r2["finalOk"]
    total = r1["rows"] + r2["rows"]
    delta_total = len(r1["ceilDeltaRows"]) + len(r2["ceilDeltaRows"])
    lines += ["## Totals", "",
              f"- Legacy golden FINAL IG: **{total_ok} / {total}**",
              f"- intermediate rows: **{r1['interOk'] + r2['interOk']} / {total}**",
              f"- rounding-strategy delta rows: **{delta_total}**",
              f"- worst float delta: **{max(r1['maxDelta'], r2['maxDelta']):.3e}**"]
    out = os.path.join(RESULTS, "GOLDEN_EXTRACTION_REPORT.md")
    with open(out, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print("\n".join(lines))


if __name__ == "__main__":
    main()
