# -*- coding: utf-8 -*-
"""Engine vs the engineers' own numbers — the decisive test for any engine change.

For each project, run the CURRENT engine (the CLI built from this checkout) over the committed
geometry fixture and the project's own workbook, and compare every conflict's FINAL IG with the value
the engineer wrote by hand in that workbook (Input Distances, column AK). Optionally give a "before"
analysis.json per project to see exactly which rows the current build moved, and in which direction.

    python scripts/prove_against_manual.py [--before-ex1 A.json] [--before-ex2 B.json]
                                           [--lin GEOMETRY.json WORKBOOK.xlsx [BEFORE.json]]

A change that moves the engine towards the human values is evidence; one that moves it away needs
an engineering argument, not a code argument (see docs/r12/LIN_R12_DEFECTS.md for the r13 example).
"""
import argparse, json, os, subprocess, sys, tempfile, warnings
import openpyxl

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CLI = os.path.join(REPO, "src", "Mahod.Intergreen.Cli", "bin", "Release", "net8.0", "Mahod.Intergreen.Cli.exe")
FIX = os.path.join(REPO, "tests", "fixtures", "geometry")
MAT = os.path.normpath(os.path.join(REPO, "..", "materials", "Inter-green Automation"))

PROJECTS = {
    "EXAMPLE1": (os.path.join(FIX, "example1.iggeometry.json"),
                 os.path.join(MAT, "03 Example 1", "05_PINES-HABAL- SHEM-TOM_IG_matrix_2026-07-28-maya.xlsx"), []),
    "EXAMPLE2": (os.path.join(FIX, "example2.iggeometry.json"),
                 os.path.join(MAT, "04 Example 2", "Abarbanel-HaMaccabim_IG_matrix_2026-07-05.xlsx"),
                 ["E-L.b1", "S-L.b2"]),
}


def manual_rows(xlsx):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb["Input Distances"]
    out = {}
    for r in range(3, ws.max_row + 1):
        cl, en, ig = ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 37).value
        if cl and en and isinstance(ig, (int, float)):
            out[f"{str(cl).strip()}→{str(en).strip()}"] = (round(float(ig)), ws.cell(r, 4).value, ws.cell(r, 5).value)
    return out


def run_engine(tag, geometry, workbook, confirmed):
    work = tempfile.mkdtemp(prefix=f"prove-{tag.lower()}-")
    geo = os.path.join(work, "g.iggeometry.json")
    with open(geometry, "rb") as s, open(geo, "wb") as d:
        d.write(s.read())
    json.dump({"schemaVersion": 1, "unitsConfirmed": "meters", "unitsConfirmedBy": "prove_against_manual",
               "confirmedEndpointReferences": confirmed, "workbook": workbook},
              open(os.path.join(work, "g.intergreen-project.json"), "w", encoding="utf-8"))
    r = subprocess.run([CLI, "ig-analyze", "--geometry", geo, "--workbook", workbook, "--name", tag,
                        "--out", os.path.join(work, "out"), "--units", "meters"],
                       capture_output=True, text=True, timeout=600)
    path = os.path.join(work, "out", tag + ".analysis.json")
    if not os.path.exists(path):
        sys.exit(f"{tag}: engine run failed\n{r.stdout[-1500:]}\n{r.stderr[-1500:]}")
    return json.load(open(path, encoding="utf-8"))


def governing(conflict):
    return next((p for p in conflict["Points"] if p["Id"] == conflict.get("DefiningPointId")), None)


def compare(tag, analysis, manual, before):
    eng = {c["Id"]: c for c in analysis["Conflicts"] if c.get("FinalIg") is not None}
    bef = {c["Id"]: c for c in (before or {}).get("Conflicts", []) if c.get("FinalIg") is not None}
    keys = [k for k in manual if k in eng]
    ok = sum(1 for k in keys if eng[k]["FinalIg"] == manual[k][0])
    line = f"=== {tag}: {len(keys)} rows comparable to the engineer's workbook — engine == manual: {ok}/{len(keys)}"
    if bef:
        ok_before = sum(1 for k in keys if k in bef and bef[k]["FinalIg"] == manual[k][0])
        line += f"   (before: {ok_before}/{len(keys)})"
    print(line)
    if bef:
        for k in sorted(keys):
            if k in bef and bef[k]["FinalIg"] != eng[k]["FinalIg"]:
                gb, ga = governing(bef[k]), governing(eng[k])
                verdict = ("→ now matches manual" if eng[k]["FinalIg"] == manual[k][0]
                           else "→ still differs" if bef[k]["FinalIg"] != manual[k][0]
                           else "→ LEFT the manual value")
                print(f"    {k:<10} manual {manual[k][0]:<3} before {bef[k]['FinalIg']:<3} after {eng[k]['FinalIg']:<3} {verdict}")
                if gb and ga:
                    print(f"               governing CD {gb['Cd']:.2f}→{ga['Cd']:.2f}  ED {gb['Ed']:.2f}→{ga['Ed']:.2f}"
                          f"  ({ga['ClearingCurveId']} / {ga['EnteringCurveId']})")
    print()
    return ok, len(keys)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--before-ex1")
    ap.add_argument("--before-ex2")
    ap.add_argument("--lin", nargs="+", metavar="PATH",
                    help="GEOMETRY.iggeometry.json WORKBOOK.xlsx [BEFORE.analysis.json]")
    a = ap.parse_args()
    if not os.path.exists(CLI):
        sys.exit("build the solution first (Release) — CLI not found: " + CLI)

    total_ok = total = 0
    befores = {"EXAMPLE1": a.before_ex1, "EXAMPLE2": a.before_ex2}
    for tag, (geo, xlsx, confirmed) in PROJECTS.items():
        before = json.load(open(befores[tag], encoding="utf-8")) if befores.get(tag) else None
        ok, n = compare(tag, run_engine(tag, geo, xlsx, confirmed), manual_rows(xlsx), before)
        total_ok += ok; total += n
    if a.lin:
        geo, xlsx = a.lin[0], a.lin[1]
        before = json.load(open(a.lin[2], encoding="utf-8")) if len(a.lin) > 2 else None
        ok, n = compare("LIN05293", run_engine("LIN05293", geo, xlsx, ["E-T.b1", "N-R.b2"]), manual_rows(xlsx), before)
        total_ok += ok; total += n
    print("=" * 72)
    print(f"TOTAL rows compared against a human-written value: {total}   engine == engineer: {total_ok}")


if __name__ == "__main__":
    main()
